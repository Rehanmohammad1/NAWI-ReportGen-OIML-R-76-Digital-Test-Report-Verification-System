import os
import io
import qrcode
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from datetime import datetime

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx_report(session_data: dict, output_path: str) -> str:
    """
    Generate an editable standardized OIML R-76 digital test report in DOCX format.
    Guaranteed content parity with PDF report.
    """
    doc = Document()

    # Set standard margins (0.5 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    is_demo = session_data.get("is_demo_data", False)
    report_num = session_data.get("report_number", "NAWI-R76-DRAFT")

    # Header
    p_gov = doc.add_paragraph()
    p_gov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_gov = p_gov.add_run("MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION\nDEPARTMENT OF LEGAL METROLOGY — GOVERNMENT OF INDIA")
    r_gov.bold = True
    r_gov.font.size = Pt(9)
    r_gov.font.color.rgb = RGBColor(71, 85, 105)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("OIML R-76 TYPE-EVALUATION TEST REPORT FOR NON-AUTOMATIC WEIGHING INSTRUMENT (NAWI)")
    r_title.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Report Certificate No: {report_num} | Date: {session_data.get('generated_at', datetime.now().strftime('%Y-%m-%d'))}")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. Laboratory Details
    h1 = doc.add_heading("1. ISSUING LABORATORY DETAILS", level=2)
    h1.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    h1.runs[0].font.size = Pt(11)

    lab = session_data.get("laboratory", {})
    t_lab = doc.add_table(rows=3, cols=4)
    t_lab.alignment = WD_TABLE_ALIGNMENT.CENTER
    lab_data = [
        [("Laboratory Name:", True), (lab.get("name", "N/A"), False), ("Accreditation Ref:", True), (lab.get("accreditation_ref", "N/A"), False)],
        [("Lab Code:", True), (lab.get("code", "N/A"), False), ("Contact Email:", True), (lab.get("contact_email", "N/A"), False)],
        [("Address:", True), (lab.get("address", "N/A"), False), ("Contact Phone:", True), (lab.get("contact_phone", "N/A"), False)]
    ]
    for r_idx, row in enumerate(lab_data):
        for c_idx, (txt, is_bold) in enumerate(row):
            cell = t_lab.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run.bold = is_bold
            run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 2. Manufacturer Details
    h2 = doc.add_heading("2. APPLICANT / MANUFACTURER DETAILS", level=2)
    h2.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    h2.runs[0].font.size = Pt(11)

    mfg = session_data.get("manufacturer", {})
    t_mfg = doc.add_table(rows=2, cols=4)
    mfg_data = [
        [("Manufacturer Name:", True), (mfg.get("name", "N/A"), False), ("Country:", True), (mfg.get("country", "India"), False)],
        [("Address:", True), (mfg.get("address", "N/A"), False), ("Contact:", True), (f"{mfg.get('contact_email','')} | {mfg.get('contact_phone','')}", False)]
    ]
    for r_idx, row in enumerate(mfg_data):
        for c_idx, (txt, is_bold) in enumerate(row):
            cell = t_mfg.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run.bold = is_bold
            run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Technical Specs
    h3 = doc.add_heading("3. INSTRUMENT IDENTIFICATION & TECHNICAL SPECIFICATIONS", level=2)
    h3.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    inst = session_data.get("instrument", {})
    t_spec = doc.add_table(rows=5, cols=4)
    spec_data = [
        [("Model Name:", True), (inst.get("model_name", "N/A"), False), ("Accuracy Class:", True), (f"Class {inst.get('accuracy_class', 'III')}", True)],
        [("Serial Number:", True), (inst.get("serial_number", "N/A"), False), ("Max Capacity (Max):", True), (f"{inst.get('max_capacity', 0)} kg", False)],
        [("Year of Manufacture:", True), (str(inst.get("year_of_manufacture", 2026)), False), ("Min Capacity (Min):", True), (f"{inst.get('min_capacity', 0)} kg", False)],
        [("Verification Interval (e):", True), (f"{inst.get('e', 0)} kg", False), ("Actual Interval (d):", True), (f"{inst.get('d', 0)} kg", False)],
        [("Scale Intervals (n):", True), (f"{inst.get('n', 0):.0f}", False), ("Temp Range:", True), (f"{inst.get('temp_min', -10)}°C to {inst.get('temp_max', 40)}°C", False)]
    ]
    for r_idx, row in enumerate(spec_data):
        for c_idx, (txt, is_bold) in enumerate(row):
            cell = t_spec.cell(r_idx, c_idx)
            set_cell_background(cell, "F1F5F9")
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run.bold = is_bold
            run.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 4. Compliance Table
    h5 = doc.add_heading("4. COMPLIANCE RULE EVALUATION AGAINST OIML R-76", level=2)
    h5.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    comp_results = session_data.get("compliance_results", [])
    t_comp = doc.add_table(rows=1 + len(comp_results), cols=6)
    headers = ["Test Procedure", "Measured Value", "Applied Limit", "Margin", "Result", "Explanation"]
    for c_idx, h_text in enumerate(headers):
        cell = t_comp.cell(0, c_idx)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        run = p.add_run(h_text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, cr in enumerate(comp_results):
        row_cells = t_comp.rows[r_idx + 1].cells
        vals = [
            (cr.get("test_procedure_code", "N/A"), True),
            (f"{cr.get('measured_value', 0):.3f} e", False),
            (str(cr.get("limit_applied", "N/A")), False),
            (f"{cr.get('margin', 0):+.3f} e", False),
            (cr.get("pass_fail", "PASS"), True),
            (cr.get("explanation_text", "N/A"), False)
        ]
        for c_idx, (txt, is_bold) in enumerate(vals):
            p = row_cells[c_idx].paragraphs[0]
            run = p.add_run(txt)
            run.bold = is_bold
            run.font.size = Pt(8.5)
            if c_idx == 4: # result badge
                if txt == "PASS":
                    run.font.color.rgb = RGBColor(22, 101, 52)
                else:
                    run.font.color.rgb = RGBColor(153, 27, 27)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 5. Determination & Authenticity
    h6 = doc.add_heading("5. OVERALL COMPLIANCE DETERMINATION", level=2)
    h6.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    p_res = doc.add_paragraph()
    p_res.add_run("OVERALL DETERMINATION RESULT: ").bold = True
    r_pf = p_res.add_run(session_data.get("overall_result", "PASS"))
    r_pf.bold = True
    r_pf.font.size = Pt(14)
    if session_data.get("overall_result") == "PASS":
        r_pf.font.color.rgb = RGBColor(22, 101, 52)
    else:
        r_pf.font.color.rgb = RGBColor(153, 27, 27)

    p_hash = doc.add_paragraph()
    r_h = p_hash.add_run(f"SHA-256 Report Hash: {session_data.get('content_hash', 'e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855')}")
    r_h.font.size = Pt(8)
    r_h.font.name = "Courier New"

    doc.save(output_path)
    return output_path
